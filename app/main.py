"""EventPilot AI — FastAPI entrypoint."""

from __future__ import annotations

import io
from pathlib import Path

from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field

from app.agent import route_user_query
from app.auth import BasicAuthMiddleware
from app.database import init_db
from app.email_service import build_confirmation_preview, email_status
from app.event_config import event_date, event_name, event_venue
from app.sample_data import SAMPLE_ROWS
from app.hackathon import build_judge_summary, export_participants_csv, parse_pass_qr_payload
from app.services import (
    bulk_send_passes,
    check_in_participant,
    confirm_participant,
    get_participant,
    list_action_logs,
    list_participants,
    replace_with_sample,
    run_tool,
    save_participants,
    send_certificate,
    send_pass,
)
from app.validation import enrich_rows

load_dotenv()

STATIC_DIR = Path(__file__).resolve().parent.parent / "static"
MAX_UPLOAD_BYTES = 2 * 1024 * 1024

app = FastAPI(title="EventPilot AI", version="1.1.0")
app.add_middleware(BasicAuthMiddleware)


@app.on_event("startup")
def on_startup() -> None:
    init_db()
    STATIC_DIR.mkdir(exist_ok=True)


class SaveRequest(BaseModel):
    rows: list[dict[str, str | None]]
    skip_invalid: bool = True


class ConfirmRequest(BaseModel):
    idempotency_key: str = Field(min_length=8, max_length=128)


class NLQueryRequest(BaseModel):
    query: str = Field(min_length=1, max_length=2000)


class CheckInRequest(BaseModel):
    qr_payload: str = Field(min_length=3, max_length=2000)


def parse_xlsx(content: bytes) -> list[dict[str, str | None]]:
    wb = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        raise HTTPException(status_code=400, detail="Workbook has no active sheet")

    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        raise HTTPException(status_code=400, detail="Empty spreadsheet")

    headers = [str(h).strip().lower() if h is not None else "" for h in header]
    name_idx = next(
        (i for i, h in enumerate(headers) if h in ("name", "full name", "participant")),
        None,
    )
    email_idx = next(
        (i for i, h in enumerate(headers) if h in ("email", "e-mail", "email address")),
        None,
    )
    id_idx = next(
        (
            i
            for i, h in enumerate(headers)
            if h in ("id", "sheet id", "sheet_id", "participant id", "participant_id", "roll", "roll no")
        ),
        None,
    )
    if name_idx is None or email_idx is None:
        raise HTTPException(
            status_code=400,
            detail="Sheet must include 'name' and 'email' columns",
        )

    out: list[dict[str, str | None]] = []
    for row in rows_iter:
        if row is None:
            continue
        name_val = row[name_idx] if name_idx < len(row) else None
        email_val = row[email_idx] if email_idx < len(row) else None
        id_val = row[id_idx] if id_idx is not None and id_idx < len(row) else None
        name = str(name_val).strip() if name_val is not None else ""
        email = str(email_val).strip() if email_val is not None else ""
        sheet_id = None
        if id_val is not None and str(id_val).strip() != "":
            if isinstance(id_val, float) and id_val.is_integer():
                sheet_id = str(int(id_val))
            else:
                sheet_id = str(id_val).strip()
        if not name and not email:
            continue
        out.append({"sheet_id": sheet_id, "name": name or None, "email": email or None})
    if not out:
        raise HTTPException(status_code=400, detail="No participant rows found")
    return out


@app.get("/")
def index():
    index_path = STATIC_DIR / "index.html"
    if not index_path.is_file():
        raise HTTPException(status_code=404, detail="Frontend not found")
    return FileResponse(index_path)


@app.get("/check-in")
def check_in_page():
    path = STATIC_DIR / "checkin.html"
    if not path.is_file():
        raise HTTPException(status_code=404, detail="Check-in page not found")
    return FileResponse(path)


app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")


@app.post("/api/upload/preview")
async def upload_preview(file: UploadFile = File(...)):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Only .xlsx files are supported")
    content = await file.read()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=400, detail="File exceeds 2 MB limit")
    rows = parse_xlsx(content)
    preview = enrich_rows(rows)
    return {
        "rows": preview,
        "total": len(preview),
        "invalid_count": sum(1 for r in preview if not r["valid"]),
    }


@app.post("/api/participants/save")
def save_body(body: SaveRequest):
    return save_participants(body.rows, skip_invalid=body.skip_invalid)


@app.get("/api/event/meta")
def event_meta():
    import os

    return {
        "name": event_name(),
        "date": event_date(),
        "venue": event_venue(),
        "organizer": os.getenv("ORGANIZER_NAME", "Event Organizer").strip() or "Event Organizer",
    }


@app.get("/api/dashboard/stats")
def dashboard_stats():
    rows = list_participants()
    total = len(rows)
    confirmed = sum(1 for r in rows if r.get("status") == "confirmed")
    pending = total - confirmed
    emailed = sum(
        1
        for r in rows
        if r.get("pass_sent_at") or r.get("certificate_sent_at") or r.get("status") == "confirmed"
    )
    delivery_rate = round(100.0 * emailed / total, 1) if total else 0.0
    conversion = round(100.0 * confirmed / total, 1) if total else 0.0
    checked_in = sum(1 for r in rows if r.get("checked_in_at"))
    attendance_rate = round(100.0 * checked_in / confirmed, 1) if confirmed else 0.0
    return {
        "total": total,
        "confirmed": confirmed,
        "pending": pending,
        "checked_in": checked_in,
        "attendance_rate": attendance_rate,
        "conversion_rate": conversion,
        "email_delivery_rate": delivery_rate,
    }


@app.get("/api/dashboard/judge-summary")
def judge_summary():
    return build_judge_summary()


@app.get("/api/participants/export.csv")
def export_csv():
    csv_text = export_participants_csv()
    return Response(
        content=csv_text,
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="eventpilot-export.csv"'},
    )


@app.post("/api/bulk/send-passes")
def bulk_passes():
    return bulk_send_passes()


@app.post("/api/check-in")
def check_in(body: CheckInRequest):
    parsed = parse_pass_qr_payload(body.qr_payload)
    if not parsed:
        raise HTTPException(status_code=400, detail="Invalid QR pass payload")
    try:
        pid = int(parsed.get("participant_id"))
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="Invalid participant id in QR")
    token = str(parsed.get("token") or "")
    result = check_in_participant(pid, token)
    if result.get("error") == "not_found":
        raise HTTPException(status_code=404, detail="Participant not found")
    if result.get("error") == "not_confirmed":
        raise HTTPException(status_code=400, detail="Registration not confirmed")
    if result.get("error") == "invalid_token":
        raise HTTPException(status_code=403, detail="Invalid or expired pass token")
    return result


@app.get("/api/participants")
def get_participants(
    search: str = Query(""),
    status: str | None = Query(None),
):
    if status and status not in ("pending", "confirmed"):
        raise HTTPException(status_code=400, detail="Invalid status filter")
    return {"participants": list_participants(search=search, status=status)}


@app.get("/api/action-logs")
def get_action_logs(limit: int = Query(100, ge=1, le=500)):
    return {"logs": list_action_logs(limit=limit)}


@app.get("/api/email/status")
def get_email_status():
    """SMTP connection status (no secrets)."""
    return email_status()


@app.get("/api/participants/{participant_id}/email-preview")
def email_preview(participant_id: int):
    p = get_participant(participant_id)
    if not p:
        raise HTTPException(status_code=404, detail="Participant not found")
    return build_confirmation_preview(p["name"], p["email"])


@app.post("/api/participants/{participant_id}/confirm")
def confirm(participant_id: int, body: ConfirmRequest):
    result = confirm_participant(participant_id, body.idempotency_key)
    if not result.get("ok"):
        raise HTTPException(status_code=404, detail="Participant not found")
    preview = None
    p = result.get("participant")
    if p:
        preview = build_confirmation_preview(p["name"], p["email"])
    return {**result, "email_preview": preview}


@app.post("/api/participants/{participant_id}/send-pass")
def send_pass_route(participant_id: int, body: ConfirmRequest):
    result = send_pass(participant_id, body.idempotency_key)
    if result.get("error") == "not_found":
        raise HTTPException(status_code=404, detail="Participant not found")
    if result.get("error") == "not_confirmed":
        raise HTTPException(status_code=400, detail="Confirm registration before sending a pass")
    return result


@app.post("/api/participants/{participant_id}/send-certificate")
def send_certificate_route(participant_id: int, body: ConfirmRequest):
    result = send_certificate(participant_id, body.idempotency_key)
    if result.get("error") == "not_found":
        raise HTTPException(status_code=404, detail="Participant not found")
    if result.get("error") == "not_confirmed":
        raise HTTPException(status_code=400, detail="Confirm registration before sending a certificate")
    return result


@app.post("/api/sample/load")
def load_sample():
    return replace_with_sample(SAMPLE_ROWS)


@app.post("/api/query")
def natural_language_query(body: NLQueryRequest):
    routed = route_user_query(body.query)
    if "error" in routed:
        return routed
    tool = routed["tool"]
    result = run_tool(tool)
    return {"routed_tool": tool, "result": result}


@app.get("/api/sample/download")
def download_sample_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Participants"
    ws.append(["name", "email"])
    for row in SAMPLE_ROWS:
        ws.append([row.get("name") or "", row.get("email") or ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="sample_participants.xlsx"'},
    )
